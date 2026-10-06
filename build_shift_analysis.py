#!/usr/bin/env python3
"""Анализ журнала проходов: рабочее место (чистая зона) и территория (турникет)."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAW = """
5:09:25	01.09.2026	Нормальный выход по ключу	Дверь в чистую зону
5:44:19	01.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:59:39	01.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:13:38	01.09.2026	Фактический выход	Турникет на проходной
20:45:29	01.09.2026	Нормальный вход по ключу	Дверь в чистую зону
22:06:51	01.09.2026	Нормальный выход по ключу	Дверь в чистую зону
22:07:59	01.09.2026	Нормальный вход по ключу	Дверь в чистую зону
22:18:37	01.09.2026	Нормальный выход по ключу	Дверь в чистую зону
22:41:15	01.09.2026	Нормальный вход по ключу	Дверь в чистую зону
7:03:48	02.09.2026	Нормальный выход по ключу	Дверь в чистую зону
7:44:49	02.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:54:35	02.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:31:02	02.09.2026	Фактический выход	Турникет на проходной
8:25:50	05.09.2026	Фактический вход	Турникет на проходной
8:46:17	05.09.2026	Нормальный вход по ключу	Дверь в чистую зону
12:14:49	05.09.2026	Нормальный выход по ключу	Дверь в чистую зону
12:30:59	05.09.2026	Нормальный вход по ключу	Дверь в чистую зону
16:07:12	05.09.2026	Нормальный выход по ключу	Дверь в чистую зону
16:39:43	05.09.2026	Нормальный вход по ключу	Дверь в чистую зону
19:09:47	05.09.2026	Нормальный выход по ключу	Дверь в чистую зону
19:42:10	05.09.2026	Нормальный вход по ключу	Дверь в чистую зону
20:57:49	05.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:10:08	05.09.2026	Фактический выход	Турникет на проходной
8:25:12	06.09.2026	Фактический вход	Турникет на проходной
8:54:03	06.09.2026	Нормальный вход по ключу	Дверь в чистую зону
9:46:14	06.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:47:04	06.09.2026	Нормальный вход по ключу	Дверь в чистую зону
12:56:10	06.09.2026	Нормальный выход по ключу	Дверь в чистую зону
13:51:29	06.09.2026	Нормальный вход по ключу	Дверь в чистую зону
18:57:09	06.09.2026	Нормальный выход по ключу	Дверь в чистую зону
19:32:14	06.09.2026	Нормальный вход по ключу	Дверь в чистую зону
21:03:28	06.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:16:32	06.09.2026	Фактический выход	Турникет на проходной
20:30:12	07.09.2026	Фактический вход	Турникет на проходной
20:34:21	07.09.2026	Нормальный вход по ключу	Дверь в Гардероб
20:54:07	07.09.2026	Нормальный вход по ключу	Дверь в чистую зону
1:21:48	08.09.2026	Нормальный выход по ключу	Дверь в чистую зону
1:34:00	08.09.2026	Нормальный вход по ключу	Дверь в чистую зону
5:12:15	08.09.2026	Нормальный выход по ключу	Дверь в чистую зону
5:14:45	08.09.2026	Нормальный вход по ключу	Дверь в чистую зону
6:16:18	08.09.2026	Нормальный выход по ключу	Дверь в чистую зону
6:47:24	08.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:55:56	08.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:09:24	08.09.2026	Фактический выход	Турникет на проходной
20:33:33	08.09.2026	Фактический вход	Турникет на проходной
20:38:00	08.09.2026	Нормальный вход по ключу	Дверь в Гардероб
21:01:15	08.09.2026	Нормальный вход по ключу	Дверь в чистую зону
3:04:43	09.09.2026	Нормальный выход по ключу	Дверь в чистую зону
3:44:44	09.09.2026	Нормальный вход по ключу	Дверь в чистую зону
7:20:09	09.09.2026	Нормальный выход по ключу	Дверь в чистую зону
7:50:09	09.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:59:53	09.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:10:02	09.09.2026	Фактический выход	Турникет на проходной
20:25:55	09.09.2026	Фактический вход	Турникет на проходной
20:50:10	09.09.2026	Нормальный вход по ключу	Дверь в чистую зону
1:46:45	10.09.2026	Нормальный выход по ключу	Дверь в чистую зону
2:34:21	10.09.2026	Нормальный вход по ключу	Дверь в чистую зону
7:24:35	10.09.2026	Нормальный выход по ключу	Дверь в чистую зону
7:55:46	10.09.2026	Нормальный вход по ключу	Дверь в чистую зону
9:00:21	10.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:13:54	10.09.2026	Фактический выход	Турникет на проходной
8:22:15	13.09.2026	Фактический вход	Турникет на проходной
8:40:47	13.09.2026	Нормальный вход по ключу	Дверь в чистую зону
12:06:07	13.09.2026	Нормальный выход по ключу	Дверь в чистую зону
13:02:17	13.09.2026	Нормальный вход по ключу	Дверь в чистую зону
14:55:13	13.09.2026	Нормальный выход по ключу	Дверь в чистую зону
14:59:00	13.09.2026	Нормальный вход по ключу	Дверь в чистую зону
18:26:04	13.09.2026	Нормальный выход по ключу	Дверь в чистую зону
18:37:13	13.09.2026	Нормальный вход по ключу	Дверь в чистую зону
18:37:34	13.09.2026	Нормальный выход по ключу	Дверь в чистую зону
18:56:32	13.09.2026	Нормальный вход по ключу	Дверь в чистую зону
20:46:20	13.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:01:29	13.09.2026	Фактический выход	Турникет на проходной
8:29:05	14.09.2026	Фактический вход	Турникет на проходной
8:51:29	14.09.2026	Нормальный вход по ключу	Дверь в чистую зону
12:57:54	14.09.2026	Нормальный выход по ключу	Дверь в чистую зону
13:18:31	14.09.2026	Нормальный вход по ключу	Дверь в чистую зону
15:24:26	14.09.2026	Нормальный выход по ключу	Дверь в чистую зону
15:52:10	14.09.2026	Нормальный вход по ключу	Дверь в чистую зону
19:41:11	14.09.2026	Нормальный выход по ключу	Дверь в чистую зону
19:57:56	14.09.2026	Нормальный вход по ключу	Дверь в чистую зону
20:59:57	14.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:14:36	14.09.2026	Фактический выход	Турникет на проходной
8:36:10	22.09.2026	Фактический вход	Турникет на проходной
8:41:08	22.09.2026	Нормальный вход по ключу	Дверь в Гардероб
8:59:24	22.09.2026	Нормальный вход по ключу	Дверь в чистую зону
13:22:06	22.09.2026	Нормальный выход по ключу	Дверь в чистую зону
14:10:46	22.09.2026	Нормальный вход по ключу	Дверь в чистую зону
16:18:15	22.09.2026	Нормальный выход по ключу	Дверь в чистую зону
17:10:11	22.09.2026	Нормальный вход по ключу	Дверь в чистую зону
17:54:10	22.09.2026	Нормальный выход по ключу	Дверь в чистую зону
17:55:14	22.09.2026	Нормальный вход по ключу	Дверь в чистую зону
18:00:02	22.09.2026	Нормальный выход по ключу	Дверь в чистую зону
18:33:35	22.09.2026	Нормальный вход по ключу	Дверь в чистую зону
21:03:33	22.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:12:30	22.09.2026	Нормальный вход по ключу	Дверь в Гардероб
21:16:44	22.09.2026	Фактический выход	Турникет на проходной
20:26:37	24.09.2026	Фактический вход	Турникет на проходной
20:43:27	24.09.2026	Нормальный вход по ключу	Дверь в чистую зону
21:33:15	24.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:34:08	24.09.2026	Нормальный вход по ключу	Дверь в чистую зону
22:58:18	24.09.2026	Нормальный выход по ключу	Дверь в чистую зону
23:08:06	24.09.2026	Нормальный вход по ключу	Дверь в чистую зону
1:43:32	25.09.2026	Нормальный выход по ключу	Дверь в чистую зону
1:45:09	25.09.2026	Нормальный вход по ключу	Дверь в чистую зону
5:11:23	25.09.2026	Нормальный выход по ключу	Дверь в чистую зону
5:56:19	25.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:58:54	25.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:14:29	25.09.2026	Фактический выход	Турникет на проходной
20:27:41	25.09.2026	Фактический вход	Турникет на проходной
20:31:50	25.09.2026	Нормальный вход по ключу	Дверь в Гардероб
20:43:16	25.09.2026	Нормальный вход по ключу	Дверь в чистую зону
22:34:49	25.09.2026	Нормальный выход по ключу	Дверь в чистую зону
22:49:14	25.09.2026	Нормальный вход по ключу	Дверь в чистую зону
1:39:18	26.09.2026	Нормальный выход по ключу	Дверь в чистую зону
1:41:19	26.09.2026	Нормальный вход по ключу	Дверь в чистую зону
5:07:28	26.09.2026	Нормальный выход по ключу	Дверь в чистую зону
5:41:09	26.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:54:20	26.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:22:06	26.09.2026	Фактический выход	Турникет на проходной
20:24:40	26.09.2026	Фактический вход	Турникет на проходной
20:29:01	26.09.2026	Нормальный вход по ключу	Дверь в Гардероб
20:43:26	26.09.2026	Нормальный вход по ключу	Дверь в чистую зону
1:57:40	27.09.2026	Нормальный выход по ключу	Дверь в чистую зону
2:18:00	27.09.2026	Нормальный вход по ключу	Дверь в чистую зону
5:06:28	27.09.2026	Нормальный выход по ключу	Дверь в чистую зону
5:17:03	27.09.2026	Нормальный вход по ключу	Дверь в чистую зону
8:56:43	27.09.2026	Нормальный выход по ключу	Дверь в чистую зону
9:03:26	27.09.2026	Нормальный вход по ключу	Дверь в Гардероб
9:08:40	27.09.2026	Фактический выход	Турникет на проходной
8:29:54	29.09.2026	Фактический вход	Турникет на проходной
8:49:45	29.09.2026	Нормальный вход по ключу	Дверь в чистую зону
12:11:30	29.09.2026	Нормальный выход по ключу	Дверь в чистую зону
13:01:18	29.09.2026	Нормальный вход по ключу	Дверь в чистую зону
16:22:02	29.09.2026	Нормальный выход по ключу	Дверь в чистую зону
16:25:26	29.09.2026	Нормальный вход по ключу	Дверь в чистую зону
19:10:14	29.09.2026	Нормальный выход по ключу	Дверь в чистую зону
19:23:31	29.09.2026	Нормальный вход по ключу	Дверь в чистую зону
20:48:13	29.09.2026	Нормальный выход по ключу	Дверь в чистую зону
20:55:34	29.09.2026	Нормальный вход по ключу	Дверь в Гардероб
21:01:09	29.09.2026	Фактический выход	Турникет на проходной
8:29:40	30.09.2026	Фактический выход	Турникет на проходной
8:29:48	30.09.2026	Фактический вход	Турникет на проходной
8:34:09	30.09.2026	Нормальный вход по ключу	Дверь в Гардероб
8:49:59	30.09.2026	Нормальный вход по ключу	Дверь в чистую зону
11:08:53	30.09.2026	Нормальный выход по ключу	Дверь в чистую зону
12:05:13	30.09.2026	Нормальный вход по ключу	Дверь в чистую зону
16:14:01	30.09.2026	Нормальный выход по ключу	Дверь в чистую зону
16:17:25	30.09.2026	Нормальный вход по ключу	Дверь в чистую зону
20:55:41	30.09.2026	Нормальный выход по ключу	Дверь в чистую зону
21:01:24	30.09.2026	Нормальный вход по ключу	Дверь в Гардероб
21:06:07	30.09.2026	Фактический выход	Турникет на проходной
"""

SHIFT_LEN = timedelta(hours=12)
THRESHOLD = timedelta(minutes=1)
NOTABLE_GAP = timedelta(minutes=10)
WD = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]

# Цвета
C_NAVY = "1F4E79"
C_TEAL = "0F6F6A"
C_RED_H = "A33B32"
C_GREEN_H = "1E7A46"
C_GOLD_H = "8A6D1B"
C_LATE = "F4C7C3"
C_LATE_FONT = "9C0006"
C_EARLY = "F8CBAD"
C_EARLY_FONT = "C65911"
C_MINOR = "FFF2CC"
C_MINOR_FONT = "806000"
C_OK = "C6EFCE"
C_OK_FONT = "006100"
C_NOTE = "FFF2CC"
C_GRAY = "F2F2F2"
C_GRAY_FONT = "666666"
C_ANOM = "E2D5F1"
C_ANOM_FONT = "5B2C6F"
C_WHITE = "FFFFFF"
C_ZEBRA = "F7F9FB"
C_INPUT = "D6EAF8"
C_SECTION = "1F4E79"

thin = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
thick_bottom = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="medium", color="1F4E79"),
)


def secs(td: timedelta | None) -> int | None:
    if td is None:
        return None
    return int(round(td.total_seconds()))


def xl(td: timedelta | None):
    if td is None:
        return None
    return td.total_seconds() / 86400.0


def clock(dt: datetime) -> str:
    return dt.strftime("%H:%M:%S")


def fmt_ru(td: timedelta | None) -> str:
    if td is None:
        return "нет данных"
    total = secs(td)
    sign = "−" if total < 0 else ""
    total = abs(total)
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    parts = []
    if h:
        parts.append(f"{h} ч")
    if m or (h and s):
        parts.append(f"{m} мин")
    if s or not parts:
        parts.append(f"{s} сек")
    return sign + " ".join(parts)


def overlap(a1, a2, b1, b2) -> timedelta:
    start = max(a1, b1)
    end = min(a2, b2)
    if end > start:
        return end - start
    return timedelta(0)


@dataclass
class Event:
    dt: datetime
    name: str
    source: str
    kind: str


@dataclass
class Interval:
    enter: datetime
    leave: datetime

    @property
    def duration(self) -> timedelta:
        return self.leave - self.enter


@dataclass
class Gap:
    start: datetime
    end: datetime
    kind: str  # absence, late, early, unknown_before

    @property
    def duration(self) -> timedelta:
        return self.end - self.start


@dataclass
class Shift:
    events: list[Event] = field(default_factory=list)
    gate_in: Event | None = None
    gate_out: Event | None = None
    opened_without_gate: bool = False
    anomalies: list[Event] = field(default_factory=list)
    start: datetime | None = None
    end: datetime | None = None
    night: bool = False
    intervals: list[Interval] = field(default_factory=list)
    gaps: list[Gap] = field(default_factory=list)
    orphan_exit: Event | None = None
    incomplete_arrival: bool = False
    first_entry: datetime | None = None
    last_exit: datetime | None = None
    workplace: timedelta = timedelta(0)
    workplace_in: timedelta = timedelta(0)
    before: timedelta = timedelta(0)
    after: timedelta = timedelta(0)
    lateness: timedelta | None = None
    early: timedelta | None = None
    breaks: timedelta | None = None
    territory: timedelta | None = None
    off_workplace_on_site: timedelta | None = None
    status: str = ""
    explanation: str = ""
    gaps_text: str = ""

    @property
    def label(self) -> str:
        assert self.start and self.end
        if self.night:
            return (
                f"{self.start:%d.%m}–{self.end:%d.%m} ночная "
                f"({WD[self.start.weekday()]}–{WD[self.end.weekday()]})"
            )
        return f"{self.start:%d.%m} дневная ({WD[self.start.weekday()]})"

    @property
    def schedule(self) -> str:
        assert self.start and self.end
        if self.night:
            return f"{self.start:%d.%m.%Y %H:%M} – {self.end:%d.%m.%Y %H:%M}"
        return f"{self.start:%d.%m.%Y} {self.start:%H:%M}–{self.end:%H:%M}"

    @property
    def significant_late(self) -> bool:
        return self.lateness is not None and self.lateness >= THRESHOLD

    @property
    def significant_early(self) -> bool:
        return self.early is not None and self.early >= THRESHOLD

    @property
    def minor_early(self) -> bool:
        return self.early is not None and timedelta(0) < self.early < THRESHOLD

    @property
    def full_workplace(self) -> bool:
        return not self.incomplete_arrival and self.first_entry is not None


def parse_events() -> list[Event]:
    events: list[Event] = []
    for line in RAW.strip().splitlines():
        time_s, date_s, name, source = [p.strip() for p in line.split("\t")]
        dt = datetime.strptime(f"{date_s} {time_s}", "%d.%m.%Y %H:%M:%S")
        if source == "Дверь в чистую зону" and "выход" in name:
            kind = "clean_out"
        elif source == "Дверь в чистую зону" and "вход" in name:
            kind = "clean_in"
        elif source == "Турникет на проходной" and "выход" in name:
            kind = "gate_out"
        elif source == "Турникет на проходной" and "вход" in name:
            kind = "gate_in"
        elif source == "Дверь в Гардероб" and "вход" in name:
            kind = "wardrobe_in"
        else:
            raise ValueError(f"Неизвестное событие: {name} / {source}")
        events.append(Event(dt, name, source, kind))
    events.sort(key=lambda e: e.dt)
    if len(events) != 152:
        raise AssertionError(f"Ожидалось 152 события, получено {len(events)}")
    for prev, cur in zip(events, events[1:]):
        if cur.dt <= prev.dt:
            raise AssertionError(f"Нарушен порядок времени: {prev.dt} >= {cur.dt}")
    return events


def group_shifts(events: list[Event]) -> list[Shift]:
    shifts: list[Shift] = []
    current: Shift | None = None
    for idx, event in enumerate(events):
        if event.kind == "gate_out" and current is None:
            nxt = events[idx + 1] if idx + 1 < len(events) else None
            if (
                nxt
                and nxt.kind == "gate_in"
                and timedelta(0) <= (nxt.dt - event.dt) <= timedelta(minutes=2)
            ):
                # Приклеим аномалию к следующей смене, когда она откроется.
                pending = getattr(group_shifts, "_pending", [])
                pending.append(event)
                group_shifts._pending = pending  # type: ignore[attr-defined]
                continue
            raise AssertionError(f"Выход с территории без открытой смены: {event.dt}")
        if event.kind == "gate_in":
            if current is not None:
                raise AssertionError(f"Новый вход на территорию до закрытия смены: {event.dt}")
            current = Shift()
            pending = getattr(group_shifts, "_pending", [])
            current.anomalies.extend(pending)
            group_shifts._pending = []  # type: ignore[attr-defined]
            current.events.append(event)
            current.gate_in = event
            continue
        if current is None:
            current = Shift(opened_without_gate=True)
            current.events.append(event)
            continue
        current.events.append(event)
        if event.kind == "gate_out":
            current.gate_out = event
            shifts.append(current)
            current = None
    if current is not None:
        raise AssertionError("Последняя смена не закрыта выходом с территории")
    if getattr(group_shifts, "_pending", []):
        raise AssertionError("Осталась непривязанная аномалия турникета")
    if len(shifts) != 15:
        raise AssertionError(f"Ожидалось 15 смен, получено {len(shifts)}")
    return shifts


def classify(shift: Shift) -> None:
    anchor = shift.gate_out.dt if shift.gate_out else shift.events[-1].dt
    if anchor.hour <= 10:
        shift.night = True
        shift.end = anchor.replace(hour=9, minute=0, second=0, microsecond=0)
    elif anchor.hour >= 21:
        shift.night = False
        shift.end = anchor.replace(hour=21, minute=0, second=0, microsecond=0)
    else:
        raise AssertionError(f"Неоднозначный конец смены: {anchor}")
    shift.start = shift.end - SHIFT_LEN
    if not (shift.start - timedelta(hours=2) <= shift.events[0].dt <= shift.end):
        raise AssertionError(
            f"Первое событие вне окна смены {shift.label}: {shift.events[0].dt}"
        )


def pair_clean_zone(shift: Shift) -> None:
    clean = [e for e in shift.events if e.kind in ("clean_in", "clean_out")]
    idx = 0
    if clean and clean[0].kind == "clean_out":
        shift.orphan_exit = clean[0]
        shift.incomplete_arrival = True
        idx = 1
    while idx < len(clean):
        if clean[idx].kind != "clean_in":
            raise AssertionError(f"{shift.label}: ожидался вход, получено {clean[idx]}")
        enter = clean[idx]
        idx += 1
        if idx >= len(clean) or clean[idx].kind != "clean_out":
            raise AssertionError(f"{shift.label}: вход {enter.dt} без выхода")
        leave = clean[idx]
        if leave.dt <= enter.dt:
            raise AssertionError(f"{shift.label}: выход не позже входа")
        shift.intervals.append(Interval(enter.dt, leave.dt))
        idx += 1
    if not shift.intervals and shift.orphan_exit is None:
        raise AssertionError(f"{shift.label}: нет событий чистой зоны")
    if shift.intervals:
        shift.first_entry = shift.intervals[0].enter
        shift.last_exit = shift.intervals[-1].leave
    elif shift.orphan_exit:
        shift.last_exit = shift.orphan_exit.dt


def measure(shift: Shift) -> None:
    assert shift.start and shift.end
    shift.workplace = sum((iv.duration for iv in shift.intervals), timedelta(0))
    far_past = shift.start - timedelta(days=2)
    far_future = shift.end + timedelta(days=2)
    shift.workplace_in = sum(
        (overlap(iv.enter, iv.leave, shift.start, shift.end) for iv in shift.intervals),
        timedelta(0),
    )
    shift.before = sum(
        (overlap(iv.enter, iv.leave, far_past, shift.start) for iv in shift.intervals),
        timedelta(0),
    )
    shift.after = sum(
        (overlap(iv.enter, iv.leave, shift.end, far_future) for iv in shift.intervals),
        timedelta(0),
    )
    if shift.incomplete_arrival or shift.first_entry is None:
        shift.lateness = None
    elif shift.first_entry > shift.start:
        shift.lateness = shift.first_entry - shift.start
    else:
        shift.lateness = timedelta(0)

    if shift.last_exit is None:
        shift.early = None
    elif shift.last_exit < shift.end:
        shift.early = shift.end - shift.last_exit
    else:
        shift.early = timedelta(0)

    if shift.lateness is not None and shift.early is not None:
        shift.breaks = SHIFT_LEN - shift.workplace_in - shift.lateness - shift.early
        if secs(shift.breaks) < 0:
            raise AssertionError(f"{shift.label}: отрицательные отлучки {shift.breaks}")
    else:
        shift.breaks = None

    # Отлучки между выходом и следующим входом, плюс ранний уход и опоздание.
    if shift.orphan_exit and shift.intervals:
        shift.gaps.append(
            Gap(shift.orphan_exit.dt, shift.intervals[0].enter, "unknown_before")
        )
    for left, right in zip(shift.intervals, shift.intervals[1:]):
        shift.gaps.append(Gap(left.leave, right.enter, "absence"))
    if shift.lateness and shift.lateness > timedelta(0) and shift.first_entry:
        shift.gaps.append(Gap(shift.start, shift.first_entry, "late"))
    if shift.early and shift.early > timedelta(0) and shift.last_exit:
        shift.gaps.append(Gap(shift.last_exit, shift.end, "early"))

    if shift.full_workplace:
        gap_inside = sum(
            (
                overlap(g.start, g.end, shift.start, shift.end)
                for g in shift.gaps
                if g.kind == "absence"
            ),
            timedelta(0),
        )
        # Хвост до первого известного входа не должен встречаться у полной смены.
        if secs(gap_inside) != secs(shift.breaks):
            raise AssertionError(
                f"{shift.label}: отлучки {shift.breaks} != интервалы {gap_inside}"
            )
        check = shift.workplace_in + shift.lateness + shift.early + shift.breaks
        if secs(check) != secs(SHIFT_LEN):
            raise AssertionError(f"{shift.label}: расклад 12 часов не сходится: {check}")

    if shift.gate_in and shift.gate_out:
        shift.territory = shift.gate_out.dt - shift.gate_in.dt
        # Всё время в чистой зоне должно лежать внутри визита на территорию.
        for iv in shift.intervals:
            if iv.enter < shift.gate_in.dt or iv.leave > shift.gate_out.dt:
                raise AssertionError(f"{shift.label}: чистая зона вне территории")
        shift.off_workplace_on_site = shift.territory - shift.workplace
        if secs(shift.off_workplace_on_site) < 0:
            raise AssertionError(f"{shift.label}: рабочее место дольше территории")
    else:
        shift.territory = None
        shift.off_workplace_on_site = None

    notable = [g for g in shift.gaps if g.kind == "absence" and g.duration >= NOTABLE_GAP]
    shift.gaps_text = "; ".join(
        f"{clock(g.start)}–{clock(g.end)} ({fmt_ru(g.duration)})" for g in notable
    )
    if not shift.gaps_text:
        shift.gaps_text = "—"


def explain(shift: Shift) -> None:
    parts: list[str] = []
    assert shift.start and shift.end and shift.last_exit

    if shift.incomplete_arrival:
        parts.append(
            "В журнале только окончание ночной смены: нет входа через турникет и неизвестно, "
            f"когда сотрудник впервые зашёл в чистую зону. Время на рабочем месте до выхода "
            f"в {clock(shift.orphan_exit.dt)} не посчитано."
        )

    if shift.opened_without_gate and not shift.incomplete_arrival:
        parts.append(
            "В журнале нет фактического входа через турникет, поэтому время на территории "
            "не посчитано."
        )

    if shift.anomalies:
        for anomaly in shift.anomalies:
            parts.append(
                f"{anomaly.dt:%d.%m.%Y} в {clock(anomaly.dt)} турникет записал фактический выход "
                "за несколько секунд до входа. Предыдущая смена уже закрыта, предшествующего входа нет. "
                "Этот выход не засчитан как уход с территории: похоже на повторное срабатывание или сбой."
            )

    if shift.significant_late and shift.first_entry:
        context = []
        if shift.gate_in:
            context.append(f"на территорию зашёл в {clock(shift.gate_in.dt)}")
        for event in shift.events:
            if event.kind == "wardrobe_in" and event.dt < shift.first_entry:
                context.append(f"в гардероб — в {clock(event.dt)}")
        context.append(f"в чистую зону — в {clock(shift.first_entry)}")
        parts.append(
            f"Опоздание на рабочее место на {fmt_ru(shift.lateness)}: "
            + ", ".join(context)
            + f" при начале смены в {shift.start:%H:%M}."
        )

    if shift.significant_early:
        sentence = (
            f"Ранний уход с рабочего места на {fmt_ru(shift.early)}: "
            f"выход из чистой зоны в {clock(shift.last_exit)} при окончании смены "
            f"в {shift.end:%H:%M}."
        )
        tail = []
        for event in shift.events:
            if event.dt <= shift.last_exit:
                continue
            if event.kind == "wardrobe_in":
                tail.append(f"вход в гардероб в {clock(event.dt)}")
            elif event.kind == "gate_out":
                tail.append(f"выход с территории в {clock(event.dt)}")
        if tail:
            sentence += " Далее " + ", ".join(tail) + "."
        if (
            not shift.significant_late
            and not shift.incomplete_arrival
            and shift.first_entry
            and shift.first_entry <= shift.start
        ):
            sentence += (
                f" К началу смены на рабочем месте уже был "
                f"(вход в чистую зону в {clock(shift.first_entry)})."
            )
        parts.append(sentence)
    elif shift.minor_early:
        parts.append(
            f"Выход из чистой зоны в {clock(shift.last_exit)} — на {fmt_ru(shift.early)} "
            f"раньше {shift.end:%H:%M}. Это меньше одной минуты и как нарушение не выделено."
        )

    if shift.significant_late or shift.significant_early or shift.incomplete_arrival or shift.anomalies or shift.minor_early or (shift.opened_without_gate and not shift.incomplete_arrival):
        shift.explanation = " ".join(parts)
    else:
        shift.explanation = "—"

    if shift.incomplete_arrival:
        shift.status = "Неполные данные"
    elif shift.significant_late and shift.significant_early:
        shift.status = "Опоздание и ранний уход"
    elif shift.significant_late:
        shift.status = "Опоздание"
    elif shift.significant_early:
        shift.status = "Ранний уход"
    else:
        shift.status = "Норма"


def analyze() -> list[Shift]:
    group_shifts._pending = []  # type: ignore[attr-defined]
    events = parse_events()
    shifts = group_shifts(events)
    for shift in shifts:
        classify(shift)
        pair_clean_zone(shift)
        measure(shift)
        explain(shift)
    # Контрольные значения, посчитанные вручную.
    by_start = {s.start: s for s in shifts}
    s05 = by_start[datetime(2026, 9, 5, 9, 0)]
    s13 = by_start[datetime(2026, 9, 13, 9, 0)]
    s0809 = by_start[datetime(2026, 9, 8, 21, 0)]
    s22 = by_start[datetime(2026, 9, 22, 9, 0)]
    s30 = by_start[datetime(2026, 9, 30, 9, 0)]
    assert secs(s05.workplace) == 10 * 3600 + 50 * 60 + 28, s05.workplace
    assert secs(s05.territory) == 12 * 3600 + 44 * 60 + 18, s05.territory
    assert secs(s05.early) == 2 * 60 + 11, s05.early
    assert secs(s13.early) == 13 * 60 + 40, s13.early
    assert secs(s0809.lateness) == 75, s0809.lateness
    assert secs(s22.workplace) == 9 * 3600 + 48 * 60 + 56, s22.workplace
    assert secs(s22.territory) == 12 * 3600 + 40 * 60 + 34, s22.territory
    assert len(s30.anomalies) == 1
    assert s30.anomalies[0].dt == datetime(2026, 9, 30, 8, 29, 40)
    return shifts


# --- Excel ---

FILL = {
    "navy": PatternFill("solid", fgColor=C_NAVY),
    "teal": PatternFill("solid", fgColor=C_TEAL),
    "red": PatternFill("solid", fgColor=C_RED_H),
    "green": PatternFill("solid", fgColor=C_GREEN_H),
    "gold": PatternFill("solid", fgColor=C_GOLD_H),
    "late": PatternFill("solid", fgColor=C_LATE),
    "early": PatternFill("solid", fgColor=C_EARLY),
    "minor": PatternFill("solid", fgColor=C_MINOR),
    "ok": PatternFill("solid", fgColor=C_OK),
    "note": PatternFill("solid", fgColor=C_NOTE),
    "gray": PatternFill("solid", fgColor=C_GRAY),
    "anom": PatternFill("solid", fgColor=C_ANOM),
    "white": PatternFill("solid", fgColor=C_WHITE),
    "input": PatternFill("solid", fgColor=C_INPUT),
    "section": PatternFill("solid", fgColor=C_SECTION),
    "total": PatternFill("solid", fgColor="E7EEF5"),
}

FONT_WHITE = Font(name="Calibri", size=11, bold=True, color=C_WHITE)
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color=C_WHITE)
FONT = Font(name="Calibri", size=11, color="1F2933")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="1F2933")
FONT_SMALL = Font(name="Calibri", size=10, color="1F2933")
FONT_TITLE = Font(name="Calibri", size=20, bold=True, color=C_NAVY)
FONT_LEAD = Font(name="Calibri", size=13, bold=True, color="1F2933")
FONT_SECTION = Font(name="Calibri", size=13, bold=True, color=C_WHITE)
FONT_LATE = Font(name="Calibri", size=11, bold=True, color=C_LATE_FONT)
FONT_EARLY = Font(name="Calibri", size=11, bold=True, color=C_EARLY_FONT)
FONT_OK = Font(name="Calibri", size=11, bold=True, color=C_OK_FONT)
FONT_MINOR = Font(name="Calibri", size=11, color=C_MINOR_FONT)
FONT_GRAY = Font(name="Calibri", size=11, italic=True, color=C_GRAY_FONT)
FONT_ANOM = Font(name="Calibri", size=11, color=C_ANOM_FONT)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

FMT_TIME = "DD.MM.YYYY HH:MM:SS"
FMT_DUR = "[h]:mm:ss"
FMT_PCT = "0.0%"


def apply_range_border(ws, r1, c1, r2, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = thin


def write_duration(cell, td: timedelta | None, *, missing_fill=True):
    cell.number_format = FMT_DUR
    cell.alignment = CENTER
    cell.font = FONT
    if td is None:
        cell.value = None
        if missing_fill:
            cell.fill = FILL["gray"]
            cell.font = FONT_GRAY
    else:
        cell.value = xl(td)


def status_font_fill(status: str):
    if status == "Норма":
        return FONT_OK, FILL["ok"]
    if status == "Опоздание":
        return FONT_LATE, FILL["late"]
    if status == "Ранний уход":
        return FONT_EARLY, FILL["early"]
    if status == "Опоздание и ранний уход":
        return FONT_LATE, FILL["late"]
    if status == "Неполные данные":
        return FONT_GRAY, FILL["gray"]
    return FONT, FILL["white"]


def build_shifts_sheet(wb: Workbook, shifts: list[Shift]) -> None:
    ws = wb.create_sheet("Смены")
    groups = [
        (1, 3, "Смена", "navy"),
        (4, 6, "Территория (турникет)", "navy"),
        (7, 8, "Рабочее место: приход и уход", "teal"),
        (9, 11, "Отклонения от графика", "red"),
        (12, 17, "Сколько был на рабочем месте", "green"),
        (18, 18, "Вне чистой зоны", "teal"),
        (19, 20, "Комментарии", "gold"),
    ]
    headers = [
        "№",
        "Смена",
        "График",
        "Вход на территорию",
        "Выход с территории",
        "Время на территории",
        "Приход на рабочее место",
        "Уход с рабочего места",
        "Опоздание",
        "Ранний уход",
        "Статус",
        "Время на рабочем месте",
        "В пределах графика",
        "До начала смены",
        "После конца смены",
        "Отлучки внутри смены",
        "Доля графика на рабочем месте",
        "На территории вне рабочего места",
        "Отлучки от 10 минут",
        "Пояснение",
    ]
    for c1, c2, title, color in groups:
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
        cell = ws.cell(1, c1, title)
        cell.fill = FILL[color]
        cell.font = FONT_HEADER
        cell.alignment = CENTER
        for c in range(c1, c2 + 1):
            ws.cell(1, c).fill = FILL[color]
            ws.cell(1, c).border = thin
            ws.cell(1, c).font = FONT_HEADER
            ws.cell(1, c).alignment = CENTER
    for col, header in enumerate(headers, 1):
        cell = ws.cell(2, col, header)
        cell.font = FONT_HEADER
        cell.alignment = CENTER
        cell.border = thin
        # Подцветка шапки по группам
        if col <= 6:
            cell.fill = FILL["navy"]
        elif col <= 8:
            cell.fill = FILL["teal"]
        elif col <= 11:
            cell.fill = FILL["red"]
        elif col <= 17:
            cell.fill = FILL["green"]
        elif col == 18:
            cell.fill = FILL["teal"]
        else:
            cell.fill = FILL["gold"]
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36

    ws.cell(2, 9).comment = Comment(
        "На сколько первый вход в чистую зону позже начала смены (09:00 или 21:00). "
        "Ноль — сотрудник уже был на рабочем месте к началу смены.",
        "Анализ",
        width=240,
        height=70,
    )
    ws.cell(2, 10).comment = Comment(
        "На сколько последний выход из чистой зоны раньше конца смены. "
        "Оранжевым выделен уход раньше чем на 1 минуту. "
        "Меньше минуты показано, но не считается нарушением.",
        "Анализ",
        width=240,
        height=80,
    )
    ws.cell(2, 17).comment = Comment(
        "Доля 12-часового графика, которую сотрудник провёл в чистой зоне. "
        "100% было бы, если бы он не выходил из чистой зоны с начала и до конца смены.",
        "Анализ",
        width=240,
        height=70,
    )

    first = 3
    for i, shift in enumerate(shifts):
        r = first + i
        values = [
            i + 1,
            shift.label + (" *" if shift.incomplete_arrival else ""),
            shift.schedule,
            shift.gate_in.dt if shift.gate_in else None,
            shift.gate_out.dt if shift.gate_out else None,
            None,  # territory duration
            shift.first_entry,
            shift.last_exit,
            None,  # late
            None,  # early
            shift.status,
            None,  # workplace
            None,
            None,
            None,
            None,
            None,
            None,
            shift.gaps_text,
            shift.explanation,
        ]
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = FONT
            cell.border = thin
            cell.alignment = CENTER if c not in (2, 3, 19, 20) else LEFT
        ws.cell(r, 4).number_format = FMT_TIME
        ws.cell(r, 5).number_format = FMT_TIME
        ws.cell(r, 7).number_format = FMT_TIME
        ws.cell(r, 8).number_format = FMT_TIME

        write_duration(ws.cell(r, 6), shift.territory)
        write_duration(ws.cell(r, 9), shift.lateness, missing_fill=shift.lateness is None)
        write_duration(ws.cell(r, 10), shift.early, missing_fill=False)
        write_duration(ws.cell(r, 12), shift.workplace, missing_fill=False)
        write_duration(ws.cell(r, 13), shift.workplace_in if shift.full_workplace else None)
        write_duration(ws.cell(r, 14), shift.before if shift.full_workplace else None)
        write_duration(ws.cell(r, 15), shift.after if shift.full_workplace else None)
        write_duration(ws.cell(r, 16), shift.breaks)
        write_duration(ws.cell(r, 18), shift.off_workplace_on_site)

        share = ws.cell(r, 17)
        share.alignment = CENTER
        share.border = thin
        share.font = FONT
        if shift.full_workplace:
            share.value = shift.workplace_in.total_seconds() / SHIFT_LEN.total_seconds()
            share.number_format = FMT_PCT
        else:
            share.value = None
            share.fill = FILL["gray"]

        if shift.gate_in is None:
            for c in (4, 6):
                ws.cell(r, c).fill = FILL["gray"]
                ws.cell(r, c).font = FONT_GRAY
            if ws.cell(r, 4).value is None:
                ws.cell(r, 4).value = "нет в журнале"

        # Подсветка опоздания и раннего ухода
        if shift.significant_late:
            for c in (7, 9):
                ws.cell(r, c).fill = FILL["late"]
                ws.cell(r, c).font = FONT_LATE
        elif shift.first_entry and shift.first_entry <= shift.start:
            ws.cell(r, 7).fill = FILL["ok"]
            ws.cell(r, 7).font = FONT_OK

        if shift.significant_early:
            for c in (8, 10):
                ws.cell(r, c).fill = FILL["early"]
                ws.cell(r, c).font = FONT_EARLY
        elif shift.minor_early:
            ws.cell(r, 10).fill = FILL["minor"]
            ws.cell(r, 10).font = FONT_MINOR
            ws.cell(r, 8).fill = FILL["ok"]
            ws.cell(r, 8).font = FONT_OK
        elif shift.early == timedelta(0):
            ws.cell(r, 8).fill = FILL["ok"]
            ws.cell(r, 8).font = FONT_OK
            ws.cell(r, 10).fill = FILL["ok"]
            ws.cell(r, 10).font = FONT_OK

        st_font, st_fill = status_font_fill(shift.status)
        ws.cell(r, 11).font = st_font
        ws.cell(r, 11).fill = st_fill

        note = ws.cell(r, 20)
        if shift.explanation != "—":
            if shift.significant_late or shift.significant_early or shift.incomplete_arrival or shift.anomalies:
                note.fill = FILL["note"]
                note.font = FONT_BOLD
            else:
                note.fill = FILL["minor"]
                note.font = FONT_MINOR
        else:
            note.font = FONT_GRAY

        text_len = len(shift.explanation) + len(shift.gaps_text)
        ws.row_dimensions[r].height = 34 if text_len < 140 else 50 if text_len < 280 else 78

    last = first + len(shifts) - 1
    if not shifts[0].incomplete_arrival:
        raise AssertionError("Первая смена должна быть неполным фрагментом 31.08–01.09")
    sum_from = first + 1
    total_row = last + 2
    ws.cell(total_row, 1, "").fill = FILL["total"]
    label = ws.cell(
        total_row,
        2,
        "Итого по 14 полным сменам. Фрагмент 31.08–01.09 в сумму не входит.",
    )
    label.font = FONT_BOLD
    label.alignment = LEFT
    label.fill = FILL["total"]
    for c in range(1, 21):
        ws.cell(total_row, c).fill = FILL["total"]
        ws.cell(total_row, c).border = thin
        ws.cell(total_row, c).font = FONT_BOLD
    for col in (6, 9, 10, 12, 13, 14, 15, 16, 18):
        cell = ws.cell(
            total_row,
            col,
            f"=SUM({get_column_letter(col)}{sum_from}:{get_column_letter(col)}{last})",
        )
        cell.number_format = FMT_DUR
        cell.font = FONT_BOLD
        cell.alignment = CENTER
        cell.fill = FILL["total"]
        cell.border = thin
    share = ws.cell(
        total_row, 17, f"=AVERAGE({get_column_letter(17)}{sum_from}:{get_column_letter(17)}{last})"
    )
    share.number_format = FMT_PCT
    share.font = FONT_BOLD
    share.alignment = CENTER
    share.fill = FILL["total"]
    share.border = thin
    ws.row_dimensions[total_row].height = 32
    fragment = shifts[0]
    footnote = ws.cell(
        total_row + 1,
        2,
        "Неполная ночь 31.08–01.09 в итог не включена: до выхода в 05:09:25 время неизвестно, "
        f"после возвращения в 05:44:19 в чистой зоне ещё {fmt_ru(fragment.workplace)}. "
        "Сумма колонки «Ранний уход» включает и уходы короче одной минуты; "
        "как нарушение на листе «Итоги» посчитаны только уходы от 1 минуты.",
    )
    footnote.font = FONT_SMALL
    footnote.alignment = LEFT
    ws.merge_cells(start_row=total_row + 1, start_column=2, end_row=total_row + 1, end_column=20)
    ws.row_dimensions[total_row + 1].height = 32

    widths = {
        1: 5, 2: 34, 3: 34, 4: 22, 5: 22, 6: 18, 7: 24, 8: 24,
        9: 14, 10: 14, 11: 20, 12: 20, 13: 18, 14: 16, 15: 18,
        16: 20, 17: 18, 18: 22, 19: 46, 20: 62,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:T{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalDpi = 300
    ws.page_setup.verticalDpi = 300
    ws.print_title_rows = "1:2"
    ws.page_setup.scale = 60
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.oddHeader.left.text = "Нахождение на рабочем месте"
    ws.oddFooter.right.text = "Стр. &P из &N"
    ws.auto_filter.ref = f"A2:T{last}"

    rule = DataBarRule(
        start_type="num",
        start_value=0,
        end_type="num",
        end_value=1,
        color=C_GREEN_H,
        showValue=True,
        minLength=None,
        maxLength=None,
    )
    ws.conditional_formatting.add(f"Q{first}:Q{last}", rule)
    ws.auto_filter.ref = f"A2:T{last}"
    return


def build_gaps_sheet(wb: Workbook, shifts: list[Shift]) -> None:
    ws = wb.create_sheet("Отлучки")
    headers = [
        "№",
        "Смена",
        "С",
        "По",
        "Длительность",
        "Что это",
        "Попадает в график смены",
    ]
    for c, header in enumerate(headers, 1):
        cell = ws.cell(1, c, header)
        cell.fill = FILL["navy"]
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = thin
    ws.row_dimensions[1].height = 24
    ws.auto_filter.ref = "A1:G1"
    row = 2
    n = 0
    kind_title = {
        "absence": "Отлучка с рабочего места",
        "late": "Опоздание: ещё не на рабочем месте",
        "early": "Ранний уход: уже не на рабочем месте",
        "unknown_before": "Был вне чистой зоны (когда зашёл до этого — неизвестно)",
    }
    for shift in shifts:
        ordered = sorted(shift.gaps, key=lambda g: g.start)
        for gap in ordered:
            n += 1
            in_schedule = overlap(gap.start, gap.end, shift.start, shift.end)
            values = [
                n,
                shift.label,
                gap.start,
                gap.end,
                xl(gap.duration),
                kind_title[gap.kind],
                xl(in_schedule) if in_schedule > timedelta(0) else xl(timedelta(0)),
            ]
            for c, value in enumerate(values, 1):
                cell = ws.cell(row, c, value)
                cell.font = FONT
                cell.border = thin
                cell.alignment = CENTER if c != 6 and c != 2 else LEFT
            ws.cell(row, 3).number_format = FMT_TIME
            ws.cell(row, 4).number_format = FMT_TIME
            ws.cell(row, 5).number_format = FMT_DUR
            ws.cell(row, 7).number_format = FMT_DUR
            if gap.kind == "late" and gap.duration >= THRESHOLD:
                for c in range(1, 8):
                    ws.cell(row, c).fill = FILL["late"]
                ws.cell(row, 6).font = FONT_LATE
            elif gap.kind == "early" and gap.duration >= THRESHOLD:
                for c in range(1, 8):
                    ws.cell(row, c).fill = FILL["early"]
                ws.cell(row, 6).font = FONT_EARLY
            elif gap.kind == "early":
                for c in range(1, 8):
                    ws.cell(row, c).fill = FILL["minor"]
                ws.cell(row, 6).font = FONT_MINOR
            elif gap.kind == "unknown_before":
                for c in range(1, 8):
                    ws.cell(row, c).fill = FILL["gray"]
            elif gap.duration >= NOTABLE_GAP:
                ws.cell(row, 5).fill = FILL["input"]
            ws.row_dimensions[row].height = 22
            row += 1
    last = row - 1
    ws.auto_filter.ref = f"A1:G{last}"
    total = last + 2
    ws.cell(total, 2, "Итого длительность перечисленных интервалов").font = FONT_BOLD
    cell = ws.cell(total, 5, f"=SUM(E2:E{last})")
    cell.number_format = FMT_DUR
    cell.font = FONT_BOLD
    cell.fill = FILL["total"]
    note = ws.cell(
        total + 2,
        2,
        "Интервалы разного вида складывать между собой как «время отсутствия» можно только внутри одной смены: "
        "опоздание, отлучки и ранний уход не пересекаются. Строка итога суммирует все строки листа, "
        "включая фрагмент неполной смены.",
    )
    note.font = FONT_SMALL
    note.alignment = LEFT
    ws.merge_cells(start_row=total + 2, start_column=2, end_row=total + 2, end_column=7)
    widths = {1: 6, 2: 36, 3: 22, 4: 22, 5: 16, 6: 62, 7: 26}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build_intervals_sheet(wb: Workbook, shifts: list[Shift]) -> None:
    ws = wb.create_sheet("Интервалы")
    headers = [
        "№",
        "Смена",
        "Вход в чистую зону",
        "Выход из чистой зоны",
        "Длительность",
        "Из них внутри графика",
        "Примечание",
    ]
    for c, header in enumerate(headers, 1):
        cell = ws.cell(1, c, header)
        cell.fill = FILL["teal"]
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = thin
    ws.row_dimensions[1].height = 24
    row = 2
    n = 0
    for shift in shifts:
        for idx, iv in enumerate(shift.intervals):
            n += 1
            inside = overlap(iv.enter, iv.leave, shift.start, shift.end)
            notes = []
            is_first = idx == 0
            is_last = idx == len(shift.intervals) - 1
            if is_first and shift.significant_late:
                notes.append("Первый вход — опоздание")
            elif is_first and shift.first_entry and shift.first_entry < shift.start:
                notes.append("Первый вход до начала смены")
            elif is_first:
                notes.append("Первый вход")
            if is_last and shift.significant_early:
                notes.append("Последний выход — ранний уход")
            elif is_last and shift.after > timedelta(0):
                notes.append("Последний выход после конца смены")
            elif is_last and shift.minor_early:
                notes.append("Последний выход меньше чем на минуту раньше конца смены")
            elif is_last:
                notes.append("Последний выход")
            if not notes:
                notes.append("Находился на рабочем месте")
            values = [
                n,
                shift.label,
                iv.enter,
                iv.leave,
                xl(iv.duration),
                xl(inside),
                "; ".join(notes),
            ]
            for c, value in enumerate(values, 1):
                cell = ws.cell(row, c, value)
                cell.font = FONT
                cell.border = thin
                cell.alignment = CENTER if c not in (2, 7) else LEFT
            ws.cell(row, 3).number_format = FMT_TIME
            ws.cell(row, 4).number_format = FMT_TIME
            ws.cell(row, 5).number_format = FMT_DUR
            ws.cell(row, 6).number_format = FMT_DUR
            if is_first and shift.significant_late:
                ws.cell(row, 3).fill = FILL["late"]
                ws.cell(row, 3).font = FONT_LATE
                ws.cell(row, 7).fill = FILL["late"]
                ws.cell(row, 7).font = FONT_LATE
            if is_last and shift.significant_early:
                ws.cell(row, 4).fill = FILL["early"]
                ws.cell(row, 4).font = FONT_EARLY
                ws.cell(row, 7).fill = FILL["early"]
                ws.cell(row, 7).font = FONT_EARLY
            elif is_last and shift.minor_early:
                ws.cell(row, 4).fill = FILL["minor"]
                ws.cell(row, 7).fill = FILL["minor"]
            ws.row_dimensions[row].height = 20
            row += 1
    last = row - 1
    ws.auto_filter.ref = f"A1:G{last}"
    total = last + 2
    ws.cell(total, 2, "Итого время на рабочем месте").font = FONT_BOLD
    cell = ws.cell(total, 5, f"=SUM(E2:E{last})")
    cell.number_format = FMT_DUR
    cell.font = FONT_BOLD
    cell.fill = FILL["total"]
    cell2 = ws.cell(total, 6, f"=SUM(F2:F{last})")
    cell2.number_format = FMT_DUR
    cell2.font = FONT_BOLD
    cell2.fill = FILL["total"]
    ws.cell(total, 4, "Сумма").font = FONT_BOLD
    ws.cell(total, 4).alignment = RIGHT
    widths = {1: 6, 2: 36, 3: 24, 4: 24, 5: 16, 6: 24, 7: 62}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def event_mark(shift: Shift, event: Event) -> str:
    if event.kind == "gate_in":
        return "Вход на территорию"
    if event.kind == "gate_out":
        return "Выход с территории"
    if event.kind == "wardrobe_in":
        return "Вход в гардероб (не рабочее место)"
    if event.kind == "clean_in" and shift.first_entry and event.dt == shift.first_entry:
        if shift.significant_late:
            return "Опоздание на рабочее место"
        if event.dt < shift.start:
            return "Приход на рабочее место до начала смены"
        return "Приход на рабочее место"
    if event.kind == "clean_out" and shift.last_exit and event.dt == shift.last_exit:
        if shift.significant_early:
            return "Ранний уход с рабочего места"
        if shift.minor_early:
            return "Уход менее чем на минуту раньше конца смены"
        if event.dt > shift.end:
            return "Уход с рабочего места после конца смены"
        return "Уход с рабочего места"
    if event.kind == "clean_in":
        return "Вернулся на рабочее место"
    if event.kind == "clean_out":
        if shift.orphan_exit and event.dt == shift.orphan_exit.dt:
            return "Выход из чистой зоны (вход до него неизвестен)"
        return "Вышел с рабочего места"
    return ""


def build_journal_sheet(wb: Workbook, shifts: list[Shift]) -> None:
    ws = wb.create_sheet("Журнал")
    headers = ["№", "Дата и время", "Событие", "Источник", "Смена", "Метка"]
    for c, header in enumerate(headers, 1):
        cell = ws.cell(1, c, header)
        cell.fill = FILL["navy"]
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = thin
    ws.row_dimensions[1].height = 24
    row = 2
    n = 0
    # Аномалия показывается в той смене, к которой привязана, перед её событиями.
    for shift in shifts:
        block = list(shift.anomalies) + list(shift.events)
        block.sort(key=lambda e: e.dt)
        for event in block:
            n += 1
            if event in shift.anomalies:
                mark = "Не использован: выход через турникет без предшествующего входа"
            else:
                mark = event_mark(shift, event)
            values = [n, event.dt, event.name, event.source, shift.label, mark]
            for c, value in enumerate(values, 1):
                cell = ws.cell(row, c, value)
                cell.font = FONT
                cell.border = thin
                cell.alignment = CENTER if c in (1, 2) else LEFT
            ws.cell(row, 2).number_format = FMT_TIME
            if event in shift.anomalies:
                for c in range(1, 7):
                    ws.cell(row, c).fill = FILL["anom"]
                ws.cell(row, 6).font = FONT_ANOM
            elif mark == "Опоздание на рабочее место":
                for c in range(1, 7):
                    ws.cell(row, c).fill = FILL["late"]
                ws.cell(row, 6).font = FONT_LATE
            elif mark == "Ранний уход с рабочего места":
                for c in range(1, 7):
                    ws.cell(row, c).fill = FILL["early"]
                ws.cell(row, 6).font = FONT_EARLY
            elif "менее чем на минуту" in mark:
                ws.cell(row, 6).fill = FILL["minor"]
                ws.cell(row, 6).font = FONT_MINOR
            elif event.kind in ("gate_in", "gate_out"):
                ws.cell(row, 4).fill = FILL["input"]
            elif event.kind == "wardrobe_in":
                ws.cell(row, 4).fill = FILL["gray"]
            ws.row_dimensions[row].height = 18
            row += 1
    last = row - 1
    if n != 152:
        raise AssertionError(f"В журнале {n} строк вместо 152")
    ws.auto_filter.ref = f"A1:F{last}"
    widths = {1: 6, 2: 22, 3: 32, 4: 32, 5: 36, 6: 68}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.page_setup.fitToHeight = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.right.text = "Журнал событий, стр. &P из &N"


def build_chart_data(wb: Workbook, shifts: list[Shift]):
    ws = wb.create_sheet("Данные графика")
    headers = ["Смена", "На рабочем месте, ч", "На территории, ч", "График, ч"]
    for c, header in enumerate(headers, 1):
        cell = ws.cell(1, c, header)
        cell.font = FONT_BOLD
        cell.fill = FILL["gray"]
    for i, shift in enumerate(shifts, 2):
        ws.cell(i, 1, shift.label.replace(" *", "") + (" *" if not shift.full_workplace else ""))
        ws.cell(i, 2, round(shift.workplace.total_seconds() / 3600, 4))
        if shift.territory is not None:
            ws.cell(i, 3, round(shift.territory.total_seconds() / 3600, 4))
        else:
            ws.cell(i, 3, None)
        ws.cell(i, 4, 12 if shift.full_workplace else None)
        for c in range(1, 5):
            ws.cell(i, c).font = FONT_SMALL
            ws.cell(i, c).border = thin
        ws.cell(i, 2).number_format = "0.00"
        ws.cell(i, 3).number_format = "0.00"
        ws.cell(i, 4).number_format = "0"
    ws.column_dimensions["A"].width = 36
    for col in "BCD":
        ws.column_dimensions[col].width = 24
    ws.sheet_state = "hidden"
    return ws


def build_summary(wb: Workbook, shifts: list[Shift], chart_ws) -> None:
    ws = wb.create_sheet("Итоги", 0)
    full = [s for s in shifts if s.full_workplace]
    with_site = [s for s in full if s.territory is not None]
    fragment = next(s for s in shifts if s.incomplete_arrival)

    workplace = sum((s.workplace for s in full), timedelta(0))
    workplace_in = sum((s.workplace_in for s in full), timedelta(0))
    territory = sum((s.territory for s in with_site), timedelta(0))
    off_site = sum((s.off_workplace_on_site for s in with_site), timedelta(0))
    late_shifts = [s for s in full if s.significant_late]
    early_shifts = [s for s in full if s.significant_early]
    minor_shifts = [s for s in shifts if s.minor_early]
    lateness = sum((s.lateness for s in late_shifts), timedelta(0))
    early = sum((s.early for s in early_shifts), timedelta(0))
    breaks = sum((s.breaks for s in full), timedelta(0))
    before = sum((s.before for s in full), timedelta(0))
    after = sum((s.after for s in full), timedelta(0))
    scheduled = SHIFT_LEN * len(full)

    ws.merge_cells("B2:G2")
    title = ws["B2"]
    title.value = "Анализ нахождения сотрудника на рабочем месте"
    title.font = FONT_TITLE
    title.alignment = LEFT

    ws.merge_cells("B3:G3")
    sub = ws["B3"]
    sub.value = (
        "Журнал СКУД за 31.08.2026–30.09.2026. Рабочее место — дверь в чистую зону. "
        "Территория — турникет на проходной. Смена длится 12 часов: с 09:00 до 21:00 или с 21:00 до 09:00."
    )
    sub.font = Font(name="Calibri", size=11, color="52606D")
    sub.alignment = LEFT
    ws.row_dimensions[3].height = 32

    ws.merge_cells("B5:G5")
    lead = ws["B5"]
    lead.value = (
        f"На рабочем месте (чистая зона) за {len(full)} полностью наблюдаемых смен: {fmt_ru(workplace)}. "
        f"На территории за {len(with_site)} смен, где известны и вход, и выход через турникет: {fmt_ru(territory)}."
    )
    lead.font = FONT_LEAD
    lead.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 48

    # Легенда
    ws.merge_cells("B6:G6")
    ws["B6"].value = (
        "Красным выделено опоздание на рабочее место, оранжевым — ранний уход раньше чем на 1 минуту. "
        "Жёлтая ячейка «Пояснение» стоит там, где есть такое отклонение или журнал неполный. "
        "Уход меньше чем на минуту раньше конца смены показан, но нарушением не считается."
    )
    ws["B6"].font = Font(name="Calibri", size=10, italic=True, color="52606D")
    ws["B6"].alignment = LEFT
    ws.row_dimensions[6].height = 32

    def section(row, text):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        cell = ws.cell(row, 2, text)
        cell.font = FONT_SECTION
        cell.fill = FILL["section"]
        cell.alignment = LEFT
        for c in range(2, 8):
            ws.cell(row, c).fill = FILL["section"]
            ws.cell(row, c).font = FONT_SECTION
        ws.row_dimensions[row].height = 24

    def header_row(row, labels):
        for c, text in enumerate(labels, 2):
            cell = ws.cell(row, c, text)
            cell.font = FONT_WHITE
            cell.fill = FILL["teal"]
            cell.alignment = CENTER
            cell.border = thin
        ws.row_dimensions[row].height = 22

    def metric(row, name, td, comment, *, fill_key=None, font=None):
        name_cell = ws.cell(row, 2, name)
        name_cell.font = FONT
        name_cell.alignment = LEFT
        name_cell.border = thin
        value = ws.cell(row, 3, xl(td))
        value.number_format = FMT_DUR
        value.font = font or FONT_BOLD
        value.alignment = CENTER
        value.border = thin
        words = ws.cell(row, 4, fmt_ru(td))
        words.font = font or FONT_BOLD
        words.alignment = CENTER
        words.border = thin
        note = ws.cell(row, 5, comment)
        note.font = FONT_SMALL
        note.alignment = LEFT
        note.border = thin
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
        ws.cell(row, 6).border = thin
        ws.cell(row, 7).border = thin
        if fill_key:
            for c in range(2, 8):
                ws.cell(row, c).fill = FILL[fill_key]
        ws.row_dimensions[row].height = 36 if len(comment) > 90 else 24

    section(8, "Общее время")
    header_row(9, ["Показатель", "Время", "Прописью", "Что входит", "", ""])
    ws.merge_cells("E9:G9")
    metric(
        10,
        "На рабочем месте",
        workplace,
        f"{len(full)} смен, у которых виден и первый вход в чистую зону, и последний выход. Это основной итог.",
    )
    metric(
        11,
        "Из них внутри графика смены",
        workplace_in,
        "То же время, но только между началом и концом смены. Минуты до начала и после конца сюда не входят.",
    )
    metric(
        12,
        "На территории",
        territory,
        f"{len(with_site)} смен с входом и выходом через турникет. Ночь 01–02.09 сюда не входит: входа через турникет в журнале нет.",
    )
    metric(
        13,
        "На территории, но не на рабочем месте",
        off_site,
        "Разница двух строк выше по одним и тем же 13 сменам: гардероб, дорога от турникета и отлучки из чистой зоны, пока человек ещё на территории.",
    )
    metric(
        14,
        "Опоздания на рабочее место",
        lateness,
        f"{len(late_shifts)} смена. Считается от начала смены до первого входа в чистую зону, если этот вход позже графика хотя бы на 1 минуту.",
        fill_key="late",
        font=FONT_LATE,
    )
    metric(
        15,
        "Ранние уходы с рабочего места",
        early,
        f"{len(early_shifts)} смен. Считается от последнего выхода из чистой зоны до конца смены, если ушёл раньше хотя бы на 1 минуту.",
        fill_key="early",
        font=FONT_EARLY,
    )
    metric(
        16,
        "Отлучки внутри смены",
        breaks,
        f"По {len(full)} полным сменам: вышел из чистой зоны и потом вернулся, не заканчивая смену. Территорию при этом не покидал.",
    )
    metric(
        17,
        "Приход раньше начала смены",
        before,
        "Сколько времени уже был в чистой зоне до официального начала. Нарушением не является.",
        fill_key="ok",
        font=FONT_OK,
    )
    metric(
        18,
        "Уход позже конца смены",
        after,
        "Сколько оставался в чистой зоне после официального конца. Нарушением не является.",
        fill_key="ok",
        font=FONT_OK,
    )

    # Проверка расклада
    check = workplace_in + lateness + early + breaks
    # minor early is inside `early` only if >= 1 min. The identity uses ALL early including seconds,
    # but `early` variable here is only significant. Add minor for the check display separately.
    minor_sum = sum((s.early for s in full if s.minor_early), timedelta(0))
    # full early including minor and zero is what completes 12h. significant + minor = all early time.
    all_early = sum((s.early for s in full), timedelta(0))
    all_late = sum((s.lateness for s in full), timedelta(0))
    identity = workplace_in + all_late + all_early + breaks
    if secs(identity) != secs(scheduled):
        raise AssertionError(f"Сводка 12 часов не сходится: {identity} vs {scheduled}")

    ws.merge_cells("B20:G20")
    ws["B20"].value = (
        f"Проверка по {len(full)} сменам: график {fmt_ru(scheduled)} = "
        f"на рабочем месте внутри графика {fmt_ru(workplace_in)} + "
        f"опоздания {fmt_ru(all_late)} + ранние уходы {fmt_ru(all_early)} "
        f"(из них меньше минуты — {fmt_ru(minor_sum)}) + отлучки {fmt_ru(breaks)}."
    )
    ws["B20"].font = Font(name="Calibri", size=10, italic=True, color="52606D")
    ws["B20"].alignment = LEFT
    ws.row_dimensions[20].height = 32

    # Среднее
    section(22, "В среднем на одну полную смену (12 часов)")
    header_row(23, ["Показатель", "На рабочем месте", "Внутри графика", "Отлучки", "Ранний уход ≥ 1 мин", "Опоздание"])
    n = len(full)
    avg_labels_row = 24
    ws.cell(avg_labels_row, 2, f"Среднее, {n} смен").font = FONT
    ws.cell(avg_labels_row, 2).alignment = LEFT
    ws.cell(avg_labels_row, 2).border = thin
    averages = [
        workplace / n,
        workplace_in / n,
        breaks / n,
        early / n,
        lateness / n,
    ]
    for c, td in enumerate(averages, 3):
        cell = ws.cell(avg_labels_row, c, xl(td))
        cell.number_format = FMT_DUR
        cell.font = FONT_BOLD
        cell.alignment = CENTER
        cell.border = thin
    ws.cell(avg_labels_row, 3).fill = FILL["ok"]
    ws.cell(avg_labels_row, 6).fill = FILL["early"]
    ws.cell(avg_labels_row, 7).fill = FILL["late"]
    # header_row wrote 6 columns (B-G). averages start at column 3, that's 5 values → columns 3..7.
    # But header has 6 labels in B-G and I put the row label in B, so averages should be C-G (5 items). Good.
    ws.row_dimensions[24].height = 24

    share = workplace_in.total_seconds() / scheduled.total_seconds()
    ws.merge_cells("B25:G25")
    ws["B25"].value = (
        f"Внутри графика сотрудник проводил на рабочем месте {share:.1%} двенадцатичасовой смены. "
        f"Остальное — отлучки из чистой зоны ({breaks.total_seconds() / scheduled.total_seconds():.1%}), "
        f"ранний уход ({all_early.total_seconds() / scheduled.total_seconds():.1%}) "
        f"и опоздание ({fmt_ru(all_late)})."
    )
    ws["B25"].font = FONT
    ws["B25"].alignment = LEFT
    ws.row_dimensions[25].height = 32

    section(27, "Меньше всего времени в чистой зоне")
    header_row(
        28,
        ["Смена", "На рабочем месте", "Внутри графика", "Отлучки", "Начало и конец", ""],
    )
    ws.merge_cells("F28:G28")
    shortest = sorted(full, key=lambda s: s.workplace)[:5]
    for i, shift in enumerate(shortest):
        r = 29 + i
        ws.cell(r, 2, shift.label).alignment = LEFT
        ws.cell(r, 3, xl(shift.workplace)).number_format = FMT_DUR
        ws.cell(r, 4, xl(shift.workplace_in)).number_format = FMT_DUR
        ws.cell(r, 5, xl(shift.breaks)).number_format = FMT_DUR
        ws.cell(r, 6, shift.status)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        st_font, st_fill = status_font_fill(shift.status)
        for c in range(2, 8):
            ws.cell(r, c).border = thin
            ws.cell(r, c).font = FONT
            ws.cell(r, c).alignment = CENTER if c > 2 else LEFT
        ws.cell(r, 3).font = FONT_BOLD
        ws.cell(r, 6).font = st_font
        ws.cell(r, 6).fill = st_fill
        ws.cell(r, 7).fill = st_fill
        ws.row_dimensions[r].height = 22
    ws.row_dimensions[34].height = 8

    # Отклонения
    section(35, "Где есть отклонение")
    dev_headers = ["Смена", "График", "Опоздание", "Ранний уход", "Пояснение", ""]
    header_row(36, dev_headers)
    ws.merge_cells("F36:G36")
    deviations = [
        s
        for s in shifts
        if s.significant_late or s.significant_early
    ]
    row = 37
    for shift in deviations:
        ws.cell(row, 2, shift.label).font = FONT
        ws.cell(row, 2).alignment = LEFT
        ws.cell(row, 3, shift.schedule).font = FONT_SMALL
        ws.cell(row, 3).alignment = LEFT
        late_cell = ws.cell(row, 4, xl(shift.lateness) if shift.lateness else xl(timedelta(0)))
        early_cell = ws.cell(row, 5, xl(shift.early) if shift.early else xl(timedelta(0)))
        late_cell.number_format = FMT_DUR
        early_cell.number_format = FMT_DUR
        note = ws.cell(row, 6, shift.explanation)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
        for c in range(2, 8):
            ws.cell(row, c).border = thin
            ws.cell(row, c).alignment = LEFT if c in (2, 3, 6, 7) else CENTER
            ws.cell(row, c).font = FONT if c not in (4, 5) else FONT_BOLD
        if shift.significant_late:
            late_cell.fill = FILL["late"]
            late_cell.font = FONT_LATE
        if shift.significant_early:
            early_cell.fill = FILL["early"]
            early_cell.font = FONT_EARLY
        elif shift.minor_early:
            early_cell.fill = FILL["minor"]
            early_cell.font = FONT_MINOR
        note.fill = FILL["note"]
        note.alignment = LEFT
        note.font = FONT_SMALL
        ws.cell(row, 7).fill = FILL["note"]
        ws.row_dimensions[row].height = 48 if len(shift.explanation) < 280 else 64
        row += 1
    dev_last = row - 1

    row += 1
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
    minor_text = "Меньше одной минуты раньше конца смены (не выделено как нарушение): " + "; ".join(
        f"{s.label.split(' (')[0]} — {fmt_ru(s.early)} (выход в {clock(s.last_exit)})" for s in minor_shifts
    )
    ws.cell(row, 2, minor_text).font = FONT_MINOR
    ws.cell(row, 2).fill = FILL["minor"]
    ws.cell(row, 2).alignment = LEFT
    for c in range(2, 8):
        ws.cell(row, c).fill = FILL["minor"]
    ws.row_dimensions[row].height = 36
    minor_row = row

    row += 2
    section(row, "Неполные и спорные записи")
    row += 1
    caveats = [
        (
            "Ночь 31.08–01.09",
            "Журнал начинается уже с выхода из чистой зоны в 05:09:25. Сколько сотрудник провёл на рабочем месте "
            f"до этого момента, неизвестно. После возвращения в 05:44:19 он был в чистой зоне ещё {fmt_ru(fragment.workplace)} "
            f"и вышел в {clock(fragment.last_exit)} — на {fmt_ru(fragment.early)} раньше 09:00. "
            "Этот фрагмент не входит в общий итог времени.",
        ),
        (
            "Ночь 01.09–02.09",
            "Время на рабочем месте посчитано полностью (вход в чистую зону в 20:45:29, выход в 08:54:35). "
            "Фактического входа через турникет в журнале нет, есть только выход в 09:31:02, "
            "поэтому время на территории за эту смену неизвестно и в итог территории не входит.",
        ),
        (
            "Утро 30.09",
            "В 08:29:40 турникет записал фактический выход, а в 08:29:48 — фактический вход. "
            "Смена 29.09 уже была закрыта выходом в 21:01:09. Утренний «выход» не используется. "
            "Время на территории 30.09 считается от входа в 08:29:48 до выхода в 21:06:07.",
        ),
        (
            "Дни без записей",
            "В выгрузке нет ни одного события за 03–04.09, 11–12.09, 15–21.09, 23.09 и 28.09, "
            "а также за начало смены 31.08. Это не записано как отсутствие: по этим датам журнал просто ничего не содержит.",
        ),
    ]
    for title_text, body in caveats:
        ws.cell(row, 2, title_text).font = FONT_BOLD
        ws.cell(row, 2).alignment = LEFT
        ws.cell(row, 2).fill = FILL["gray"]
        ws.cell(row, 2).border = thin
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        body_cell = ws.cell(row, 3, body)
        body_cell.font = FONT_SMALL
        body_cell.alignment = LEFT
        body_cell.border = thin
        for c in range(3, 8):
            ws.cell(row, c).border = thin
            ws.cell(row, c).fill = FILL["white"]
        ws.row_dimensions[row].height = 48
        row += 1

    chart_row = row + 2
    ws.merge_cells(start_row=chart_row, start_column=2, end_row=chart_row, end_column=7)
    cap = ws.cell(chart_row, 2, "Часы по сменам. Звёздочка у 31.08–01.09 — неполная смена, в общий итог не входит. Пустой столбец территории — нет входа через турникет.")
    cap.font = FONT_BOLD
    cap.alignment = LEFT

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Время на рабочем месте и на территории"
    chart.y_axis.title = "Часы"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 16
    chart.y_axis.majorUnit = 2
    data = Reference(chart_ws, min_col=2, max_col=4, min_row=1, max_row=1 + len(shifts))
    cats = Reference(chart_ws, min_col=1, min_row=2, max_row=1 + len(shifts))
    chart.add_data(data, from_rows=False, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend.position = "b"
    chart.y_axis.majorGridlines = None
    chart.style = 10
    chart.overlap = -8
    chart.gapWidth = 60
    # Цвета серий: зелёный, синий, серая линия плана — план оставим столбцом нейтральным.
    colors = ["1E7A46", "2E75B6", "BFBFBF"]
    for series, color in zip(chart.series, colors):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
    chart.y_axis.numFmt = "0"
    chart.x_axis.txPr = None
    chart.width = 28
    chart.height = 10
    ws.add_chart(chart, f"B{chart_row + 1}")

    # Ширины и оформление листа
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 36
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 3
    ws.row_dimensions[2].height = 28
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = "Красный — опоздание, оранжевый — ранний уход от 1 минуты"
    ws.oddFooter.right.text = "Стр. &P из &N"
    ws.sheet_properties.tabColor = C_NAVY

    # Чтобы не ругался линтер на неиспользуемые локальные имена, если я их оставил.
    _ = (check, dev_last, minor_row)
    return {
        "full": len(full),
        "with_site": len(with_site),
        "workplace": workplace,
        "workplace_in": workplace_in,
        "territory": territory,
        "off_site": off_site,
        "lateness": lateness,
        "early": early,
        "breaks": breaks,
        "before": before,
        "after": after,
        "late_n": len(late_shifts),
        "early_n": len(early_shifts),
        "minor": minor_shifts,
        "share": share,
        "fragment": fragment,
        "scheduled": scheduled,
        "all_early": all_early,
        "deviations": deviations,
    }


def build_method(wb: Workbook) -> None:
    ws = wb.create_sheet("Методика")
    ws.sheet_properties.tabColor = "8A6D1B"
    paragraphs = [
        ("Как посчитано", True),
        (
            "Рабочее место — зона за дверью «Дверь в чистую зону». "
            "Время на рабочем месте — сумма промежутков от «Нормальный вход по ключу» до следующего "
            "«Нормальный выход по ключу» по этой двери. Короткий заход тоже учтён: 13.09 сотрудник "
            "зашёл в 18:37:13 и вышел в 18:37:34, это 21 секунда.",
            False,
        ),
        (
            "Территория — от «Фактический вход» до «Фактический выход» на турникете проходной. "
            "Гардероб стоит на территории, но рабочим местом не является. В журнале у гардероба есть только входы, "
            "выходов нет, поэтому отдельное время гардероба не считается. Оно входит во время на территории "
            "и не входит во время на рабочем месте.",
            False,
        ),
        (
            "График смены — 12 часов. Если сотрудник уходит с территории утром (около 09:00), это конец ночной смены "
            "21:00–09:00. Если уходит вечером (около 21:00), это конец дневной смены 09:00–21:00. "
            "В журнале других вариантов выхода нет.",
            False,
        ),
        (
            "Опоздание — первый вход в чистую зону позже начала смены. Приход раньше начала не считается нарушением: "
            "к 09:00 или к 21:00 человек уже был на рабочем месте.",
            False,
        ),
        (
            "Ранний уход — последний выход из чистой зоны раньше конца смены. "
            "Если после этого он ещё заходил в гардероб или выходил через турникет позже, смена на рабочем месте "
            "всё равно закончилась в момент выхода из чистой зоны. Уход позже конца смены нарушением не считается.",
            False,
        ),
        (
            "Цветом выделены отклонения от 1 минуты. Красный — опоздание, оранжевый — ранний уход. "
            "Жёлтым отмечено пояснение у такого отклонения и у неполных данных. "
            "Расхождение меньше минуты видно в цифрах и в бледно-жёлтой ячейке, но в статус «нарушение» не попадает: "
            "это сопоставимо с секундами срабатывания двери.",
            False,
        ),
        (
            "Отлучка — выход из чистой зоны и следующий вход в неё в той же смене. "
            "Повторного прохода через турникет посреди смены нет, то есть во время отлучки сотрудник оставался на территории. "
            "Отлучка не названа ни опозданием, ни ранним уходом. В колонке «Отлучки от 10 минут» перечислены только "
            "промежутки от 10 минут; более короткие лежат на листе «Отлучки».",
            False,
        ),
        (
            "Доля графика — время в чистой зоне между началом и концом смены, делённое на 12 часов. "
            "Минуты, проведённые в чистой зоне до начала или после конца, в долю не входят, но входят в общее время на рабочем месте.",
            False,
        ),
        ("Что сознательно не додумано", True),
        (
            "Дни без единой записи не названы прогулами или выходными. В файле их нет, и по одному только журналу "
            "нельзя понять, была это смена другого человека, отгул или обрезанная выгрузка.",
            False,
        ),
        (
            "Ночь 31.08–01.09 обрезана началом файла. Известен только кусок с 05:44:19 до 08:59:39 плюс сам факт, "
            "что в 05:09:25 человек уже выходил из чистой зоны. Общий итог на этот кусок не опирается.",
            False,
        ),
        (
            "Ночь 01.09–02.09 не имеет входа через турникет. В итог территории она не включена, в итог рабочего места включена.",
            False,
        ),
        (
            "Событие 30.09 в 08:29:40 («фактический выход» за 8 секунд до входа) не принято за уход с территории.",
            False,
        ),
    ]
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 140
    row = 2
    for text, is_header in paragraphs:
        cell = ws.cell(row, 2, text)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if is_header:
            cell.font = FONT_SECTION
            cell.fill = FILL["section"]
            ws.row_dimensions[row].height = 24
        else:
            cell.font = Font(name="Calibri", size=12, color="1F2933")
            ws.row_dimensions[row].height = 48 if len(text) < 320 else 64
        row += 1
    ws.row_dimensions[1].height = 10
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.6
    ws.page_margins.right = 0.6


def print_report(shifts: list[Shift], summary: dict) -> None:
    print("=" * 80)
    for s in shifts:
        print(
            f"{s.label:32} статус={s.status:20} "
            f"РМ={fmt_ru(s.workplace):18} тер={fmt_ru(s.territory):18} "
            f"опозд={fmt_ru(s.lateness):16} ранний={fmt_ru(s.early):16} "
            f"отлучки={fmt_ru(s.breaks)}"
        )
        if s.explanation != "—":
            print("   ", s.explanation)
        if s.gaps_text != "—":
            print("    отлучки:", s.gaps_text)
    print("=" * 80)
    print("РМ", fmt_ru(summary["workplace"]))
    print("РМ в графике", fmt_ru(summary["workplace_in"]))
    print("Территория", fmt_ru(summary["territory"]))
    print("Вне РМ на территории", fmt_ru(summary["off_site"]))
    print("Опоздания", summary["late_n"], fmt_ru(summary["lateness"]))
    print("Ранние", summary["early_n"], fmt_ru(summary["early"]))
    print("Отлучки", fmt_ru(summary["breaks"]))
    print("До начала", fmt_ru(summary["before"]))
    print("После конца", fmt_ru(summary["after"]))
    print("Доля", f"{summary['share']:.1%}")
    print("Фрагмент", fmt_ru(summary["fragment"].workplace))


def main() -> None:
    shifts = analyze()
    wb = Workbook()
    # Удаляем стандартный лист после того, как создадим свои: create_sheet и summary вставит Итоги на 0.
    default = wb.active
    chart_ws = build_chart_data(wb, shifts)
    summary = build_summary(wb, shifts, chart_ws)
    build_shifts_sheet(wb, shifts)
    build_gaps_sheet(wb, shifts)
    build_intervals_sheet(wb, shifts)
    build_journal_sheet(wb, shifts)
    build_method(wb)
    wb.remove(default)
    # Порядок: Итоги, Смены, Отлучки, Интервалы, Журнал, Методика. Данные графика скрыты.
    order = ["Итоги", "Смены", "Отлучки", "Интервалы", "Журнал", "Методика", "Данные графика"]
    for idx, name in enumerate(order):
        wb.move_sheet(name, offset=idx - wb.sheetnames.index(name))
    wb.properties.title = "Нахождение на рабочем месте, сентябрь 2026"
    wb.properties.creator = "Анализ журнала СКУД"
    wb.properties.subject = "Рабочее место — чистая зона; территория — турникет"
    path = "/workspace/analiz_rabochee_mesto_sentyabr_2026.xlsx"
    wb.save(path)
    print_report(shifts, summary)
    print("SAVED", path)


if __name__ == "__main__":
    main()
